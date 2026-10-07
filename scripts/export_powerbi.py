"""Export a Power BI kit from the cleaned transactions.

Writes to powerbi/:
  data/*.csv                 one file per table (the PBIP semantic model loads these)
  data/RetailAnalytics.xlsx  every table as a sheet, for "upload a file" in the Power BI service
  measures.dax               every measure, ready to paste into the browser's model editor
  measures_query.dax         Desktop DAX query view: defines all measures at once + tie-out query
  RetailAnalytics.pbip + .SemanticModel/ + .Report/   Power BI Project (generated, see scripts/pbip.py)
  BUILD_GUIDE.md             the tie-out section is refreshed with numbers computed here

Usage:
    python scripts/export_powerbi.py            # everything
    python scripts/export_powerbi.py --no-xlsx  # skip the slow Excel workbook (a few minutes)

All derived tables come from retail/metrics.py, so Power BI and the Streamlit app agree.
"""

from __future__ import annotations

import argparse
import re
import sys
import warnings
from dataclasses import dataclass, field
from datetime import date, timedelta
from pathlib import Path

import numpy as np
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from retail import metrics  # noqa: E402
from retail.countries import iso3, region  # noqa: E402

PBI = ROOT / "powerbi"
DATA = PBI / "data"
EXCEL_MAX_ROWS = 1_048_576 - 1  # minus the header row
BASKET_MIN_SUPPORT = 0.01
RETURNS_MIN_UNITS = 100


# ---------------------------------------------------------------- model description (shared with pbip.py)

@dataclass
class Column:
    name: str
    dtype: str  # TMDL dataType: string | int64 | double | dateTime | boolean
    fmt: str | None = None
    summarize: str = "none"  # none | sum
    sort_by: str | None = None
    key: bool = False
    hidden: bool = False
    category: str | None = None  # dataCategory, e.g. Country


@dataclass
class Measure:
    name: str
    table: str
    expression: str
    fmt: str
    folder: str
    description: str


@dataclass
class Table:
    name: str
    columns: list[Column]
    description: str
    is_date_table: bool = False
    frame: pd.DataFrame | None = field(default=None, repr=False)


GBP = "\\£#,0"
GBP2 = "\\£#,0.00"
INT = "#,0"
PCT = "0.0%"


def model_tables() -> list[Table]:
    return [
        Table("fact_sales", [
            Column("invoice", "string"),
            Column("invoice_date", "dateTime", fmt="yyyy-mm-dd"),
            Column("hour", "int64"),
            Column("stock_code", "string"),
            Column("customer_id", "int64", fmt="0"),
            Column("country", "string"),
            Column("quantity", "int64", fmt=INT, summarize="sum"),
            Column("price", "double", fmt=GBP2),
            Column("revenue", "double", fmt=GBP2, summarize="sum"),
            Column("is_cancellation", "boolean"),
        ], "One row per product invoice line, sales and cancellations (invoice starts with C, negative quantity)."),
        Table("dim_date", [
            Column("date", "dateTime", fmt="yyyy-mm-dd", key=True),
            Column("year", "int64", fmt="0"),
            Column("quarter", "string"),
            Column("month_num", "int64", hidden=True),
            Column("month_name", "string", sort_by="month_num"),
            Column("month_start", "dateTime", fmt="mmm yyyy"),
            Column("year_month", "string"),
            Column("week", "int64"),
            Column("weekday_num", "int64", hidden=True),
            Column("weekday_name", "string", sort_by="weekday_num"),
            Column("is_partial_month", "boolean"),
        ], "Continuous calendar, Dec 2009 - Dec 2011. Marked as the date table.", is_date_table=True),
        Table("dim_product", [
            Column("stock_code", "string", key=True),
            Column("description", "string"),
            Column("revenue_rank", "int64"),
            Column("abc_class", "string"),
        ], "Products. ABC class: A = products making the first 80% of revenue, B = next 15%, C = the rest."),
        Table("dim_customer", [
            Column("customer_id", "int64", fmt="0", key=True),
            Column("country", "string"),
            Column("acquisition_month", "dateTime", fmt="mmm yyyy"),
            Column("recency_days", "int64"),
            Column("frequency", "int64"),
            Column("monetary", "double", fmt=GBP),
            Column("r_score", "int64"),
            Column("f_score", "int64"),
            Column("m_score", "int64"),
            Column("segment_order", "int64", hidden=True),
            Column("segment", "string", sort_by="segment_order"),
        ], "Identified customers with RFM scores as of 10 Dec 2011 (fixed; does not react to date slicers)."),
        Table("dim_country", [
            Column("country", "string", key=True, category="Country"),
            Column("iso3", "string"),
            Column("region", "string"),
        ], "Countries with ISO-3 code and region."),
        Table("cohort_retention", [
            Column("cohort", "dateTime", fmt="mmm yyyy"),
            Column("cohort_label", "string"),
            Column("period", "int64"),
            Column("customers", "int64", fmt=INT),
            Column("cohort_size", "int64", fmt=INT),
            Column("retention", "double", fmt=PCT),
            Column("revenue", "double", fmt=GBP),
        ], "Monthly acquisition cohorts x months since first order (full date range)."),
        Table("basket_rules", [
            Column("antecedent", "string"),
            Column("antecedent_desc", "string"),
            Column("consequent", "string"),
            Column("consequent_desc", "string"),
            Column("pair_baskets", "int64", fmt=INT),
            Column("support", "double", fmt="0.00%"),
            Column("confidence", "double", fmt=PCT),
            Column("lift", "double", fmt="0.0"),
        ], f"Pairwise rules A -> B at {BASKET_MIN_SUPPORT:.0%} minimum support over all orders."),
        Table("product_returns", [
            Column("stock_code", "string"),
            Column("description", "string"),
            Column("units_sold", "int64", fmt=INT),
            Column("revenue", "double", fmt=GBP),
            Column("units_cancelled", "int64", fmt=INT),
            Column("cancelled_value", "double", fmt=GBP),
            Column("cancellations", "int64", fmt=INT),
            Column("return_rate", "double", fmt=PCT),
        ], f"Cancelled units and value per product (products with {RETURNS_MIN_UNITS}+ units sold)."),
    ]


# (from table, from column, to table, to column); all many-to-one, single direction
RELATIONSHIPS = [
    ("fact_sales", "invoice_date", "dim_date", "date"),
    ("fact_sales", "stock_code", "dim_product", "stock_code"),
    ("fact_sales", "customer_id", "dim_customer", "customer_id"),
    ("fact_sales", "country", "dim_country", "country"),
    ("product_returns", "stock_code", "dim_product", "stock_code"),
]

SALE = "fact_sales[is_cancellation] = FALSE()"
CANCEL = "fact_sales[is_cancellation] = TRUE()"
MEASURES = [
    Measure("Revenue", "fact_sales", "SUM(fact_sales[revenue])", GBP, "Sales",
            "Net revenue in GBP: product sales minus cancelled product lines (cancellations are negative)."),
    Measure("Gross Sales", "fact_sales", f"CALCULATE(SUM(fact_sales[revenue]), {SALE})", GBP, "Sales",
            "Product sales before cancellations."),
    Measure("Orders", "fact_sales", f"CALCULATE(DISTINCTCOUNT(fact_sales[invoice]), {SALE})", INT, "Sales",
            "Distinct sales invoices."),
    Measure("Customers", "fact_sales", f"CALCULATE(DISTINCTCOUNTNOBLANK(fact_sales[customer_id]), {SALE})", INT,
            "Sales", "Distinct identified customers; guest lines have no customer ID."),
    Measure("Avg Order Value", "fact_sales", "DIVIDE([Revenue], [Orders])", GBP2, "Sales", "Revenue / Orders."),
    Measure("Units", "fact_sales", "SUM(fact_sales[quantity])", INT, "Sales", "Units sold minus units cancelled."),
    Measure("Country Revenue Share", "fact_sales",
            "DIVIDE([Revenue], CALCULATE([Revenue], ALLSELECTED(dim_country[country])))", PCT, "Sales",
            "Share of the selected countries' revenue."),
    Measure("Revenue PY", "fact_sales", "CALCULATE([Revenue], SAMEPERIODLASTYEAR(dim_date[date]))", GBP, "Time",
            "Revenue in the same period one year earlier."),
    Measure("Revenue YoY %", "fact_sales",
            "VAR py = [Revenue PY]\nRETURN IF(NOT ISBLANK(py), DIVIDE([Revenue] - py, py))", PCT, "Time",
            "Change vs the same period last year; blank when there is no prior-year data."),
    Measure("Pareto Cumulative %", "fact_sales",
            "VAR thisRevenue = [Revenue]\n"
            "VAR ranked = ADDCOLUMNS(ALLSELECTED(dim_product[stock_code]), \"@rev\", [Revenue])\n"
            "VAR running = SUMX(FILTER(ranked, [@rev] >= thisRevenue), [@rev])\n"
            "RETURN DIVIDE(running, SUMX(ranked, [@rev]))", PCT, "Products",
            "Cumulative share of revenue from products earning at least as much as this one (80/20 curve)."),
    Measure("Cancelled Value", "fact_sales", f"CALCULATE(-SUM(fact_sales[revenue]), {CANCEL})", GBP, "Returns",
            "Value of cancelled product lines, as a positive number."),
    Measure("Cancellation Rate", "fact_sales", "DIVIDE([Cancelled Value], [Gross Sales])", PCT, "Returns",
            "Cancelled value / gross sales."),
    Measure("Units Cancelled", "fact_sales", f"CALCULATE(-SUM(fact_sales[quantity]), {CANCEL})", INT, "Returns",
            "Units on cancelled lines, as a positive number."),
    Measure("Segment Customers", "dim_customer", "COUNTROWS(FILTER(dim_customer, NOT ISBLANK(dim_customer[segment])))",
            INT, "Customers", "Customers with an RFM segment."),
    Measure("Segment Revenue", "dim_customer", "SUM(dim_customer[monetary])", GBP, "Customers",
            "Lifetime revenue of the customers in view."),
    Measure("Segment Customer Share", "dim_customer",
            "DIVIDE([Segment Customers], CALCULATE([Segment Customers], ALL(dim_customer[segment], dim_customer[segment_order])))",
            PCT, "Customers", "Share of all segmented customers."),
    Measure("Segment Revenue Share", "dim_customer",
            "DIVIDE([Segment Revenue], CALCULATE([Segment Revenue], ALL(dim_customer[segment], dim_customer[segment_order])))",
            PCT, "Customers", "Share of all segmented customers' revenue."),
    Measure("Retention %", "cohort_retention",
            "DIVIDE(SUM(cohort_retention[customers]), SUM(cohort_retention[cohort_size]))", PCT, "Cohorts",
            "Returning customers / cohort size. Across cohorts this is the size-weighted average."),
]


# ---------------------------------------------------------------- table builders

def build_frames(df: pd.DataFrame) -> dict[str, pd.DataFrame]:
    end = df["invoice_date"].max().date()
    snapshot = end + timedelta(days=1)
    products = df[df["is_product"]]

    fact = pd.DataFrame({
        "invoice": products["invoice"].astype(str),
        "invoice_date": products["invoice_date"].dt.normalize(),
        "hour": products["invoice_date"].dt.hour.astype(int),
        "stock_code": products["stock_code"].astype(str),
        "customer_id": products["customer_id"],
        "country": products["country"].astype(str),
        "quantity": products["quantity"].astype(int),
        "price": products["price"],
        "revenue": products["revenue"],
        "is_cancellation": products["is_cancellation"],
    }).reset_index(drop=True)

    days = pd.date_range(date(fact["invoice_date"].min().year, fact["invoice_date"].min().month, 1),
                         date(end.year, 12, 31), freq="D")
    data_end = pd.Timestamp(end)
    dim_date = pd.DataFrame({"date": days})
    dim_date["year"] = days.year
    dim_date["quarter"] = "Q" + days.quarter.astype(str)
    dim_date["month_num"] = days.month
    dim_date["month_name"] = days.strftime("%b")
    dim_date["month_start"] = days.to_period("M").to_timestamp()
    dim_date["year_month"] = days.strftime("%Y-%m")
    dim_date["week"] = days.isocalendar().week.astype(int).to_numpy()
    dim_date["weekday_num"] = days.dayofweek + 1
    dim_date["weekday_name"] = days.strftime("%a")
    dim_date["is_partial_month"] = (dim_date["month_start"] == data_end.to_period("M").to_timestamp()) & \
        (data_end < data_end + pd.offsets.MonthEnd(0))

    summary = metrics.product_summary(df)
    summary["revenue_rank"] = np.arange(1, len(summary) + 1)
    cum = summary["revenue"].cumsum() / summary["revenue"].sum()
    summary["abc_class"] = np.select([cum.shift(fill_value=0) < 0.80, cum.shift(fill_value=0) < 0.95], ["A", "B"], "C")
    # Every code in the fact table, including products that only appear on cancellations.
    desc = products[["stock_code", "description"]].astype(str).drop_duplicates("stock_code")
    dim_product = desc.merge(
        summary[["stock_code", "revenue_rank", "abc_class"]], on="stock_code", how="left")
    dim_product["abc_class"] = dim_product["abc_class"].fillna("No sales")
    dim_product = dim_product.sort_values(["revenue_rank", "stock_code"], na_position="last", ignore_index=True)

    rfm = metrics.rfm(df, snapshot)
    acquired = metrics.first_purchase_month(df).rename("acquisition_month")
    seg_order = {s: i for i, s in enumerate(metrics.SEGMENTS, start=1)}
    ids = pd.Series(fact["customer_id"].dropna().unique(), name="customer_id").astype("Int64")
    home = (fact.dropna(subset=["customer_id"]).groupby("customer_id")["country"]
            .agg(lambda c: c.mode().iat[0]).rename("home_country"))
    dim_customer = (
        pd.DataFrame({"customer_id": ids})
        .merge(rfm[["customer_id", "country", "recency", "frequency", "monetary", "r_score", "f_score", "m_score",
                    "segment"]], on="customer_id", how="left")
        .merge(acquired, left_on="customer_id", right_index=True, how="left")
        .merge(home, left_on="customer_id", right_index=True, how="left")
        .rename(columns={"recency": "recency_days"})
    )
    dim_customer["country"] = dim_customer["country"].fillna(dim_customer["home_country"])
    dim_customer["segment_order"] = dim_customer["segment"].map(seg_order)
    dim_customer = dim_customer.drop(columns="home_country").sort_values("customer_id", ignore_index=True)
    for col in ["recency_days", "frequency", "r_score", "f_score", "m_score", "segment_order"]:
        dim_customer[col] = dim_customer[col].astype("Int64")

    countries = sorted(fact["country"].unique())
    dim_country = pd.DataFrame({"country": countries, "iso3": [iso3(c) for c in countries],
                                "region": [region(c) for c in countries]})

    cohorts = metrics.cohort_table(df, metrics.first_purchase_month(df))
    cohorts.insert(1, "cohort_label", cohorts["cohort"].dt.strftime("%Y-%m"))

    rules = metrics.basket_rules(df, min_support=BASKET_MIN_SUPPORT)
    returns = metrics.product_returns(df, min_units_sold=RETURNS_MIN_UNITS)

    frames = {
        "fact_sales": fact,
        "dim_date": dim_date,
        "dim_product": dim_product,
        "dim_customer": dim_customer,
        "dim_country": dim_country,
        "cohort_retention": cohorts,
        "basket_rules": rules,
        "product_returns": returns[["stock_code", "description", "units_sold", "revenue", "units_cancelled",
                                    "cancelled_value", "cancellations", "return_rate"]],
    }
    # Column order follows the model definition, so CSV headers, sheets and TMDL all line up.
    return {t.name: frames[t.name][[c.name for c in t.columns]].reset_index(drop=True) for t in model_tables()}


# ---------------------------------------------------------------- writers

def write_csvs(frames: dict[str, pd.DataFrame]) -> None:
    DATA.mkdir(parents=True, exist_ok=True)
    for name, frame in frames.items():
        out = frame.copy()
        for col in out.columns:
            if pd.api.types.is_bool_dtype(out[col]):
                out[col] = out[col].map({True: "true", False: "false"})
            elif pd.api.types.is_datetime64_any_dtype(out[col]):
                out[col] = out[col].dt.strftime("%Y-%m-%d")
        out.to_csv(DATA / f"{name}.csv", index=False, float_format="%.6g" if name == "basket_rules" else None)
        print(f"  {name}.csv  {len(frame):>9,} rows")


def write_excel(frames: dict[str, pd.DataFrame]) -> Path:
    """One sheet per table, each formatted as an Excel table named after it.

    The Power BI service only imports ranges formatted as Excel tables ("We couldn't find any data
    formatted as a table" otherwise), and the table names become the model's table names, which
    the DAX measures refer to. openpyxl's write-only mode streams rows (flat RAM for the 1M-row fact
    sheet) and, unlike xlsxwriter's constant_memory mode, still supports tables.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table as XlTable, TableColumn, TableStyleInfo

    path = DATA / "RetailAnalytics.xlsx"
    if len(frames["fact_sales"]) > EXCEL_MAX_ROWS:
        raise SystemExit(f"fact_sales has {len(frames['fact_sales']):,} rows; an Excel sheet holds {EXCEL_MAX_ROWS:,}.")
    wb = Workbook(write_only=True)
    for name, frame in frames.items():
        ws = wb.create_sheet(name)
        cols = list(frame.columns)
        ref = f"A1:{get_column_letter(len(cols))}{len(frame) + 1}"
        table = XlTable(displayName=name, ref=ref,
                        tableStyleInfo=TableStyleInfo(name="TableStyleLight1", showRowStripes=True))
        table.tableColumns = [TableColumn(id=i + 1, name=c) for i, c in enumerate(cols)]
        with warnings.catch_warnings():  # openpyxl warns in write-only mode even when columns are set
            warnings.simplefilter("ignore", UserWarning)
            ws.add_table(table)
        ws.append(cols)
        values = []
        for col in cols:
            series = frame[col]
            if pd.api.types.is_datetime64_any_dtype(series):
                values.append([None if pd.isna(v) else v.date() for v in series])  # dates, not datetimes
            elif pd.api.types.is_bool_dtype(series):
                values.append(series.tolist())
            elif pd.api.types.is_numeric_dtype(series):
                values.append([None if pd.isna(v) else v for v in series.astype(object)])
            else:
                values.append([None if pd.isna(v) else str(v) for v in series])  # codes like 489434 stay text
        for row in zip(*values):
            ws.append(row)
        print(f"  sheet + table {name}: {len(frame):,} rows ({ref})")
    wb.save(path)
    print(f"  {path.relative_to(ROOT)} ({path.stat().st_size / 1e6:.1f} MB)")
    return path


def write_measures_dax() -> None:
    lines = ["// DAX measures for the Retail Analytics model. Generated by scripts/export_powerbi.py.",
             "// In the browser: open the semantic model > Open data model > select the home table > New measure,",
             "// then paste one block at a time (name = expression). Set the format shown in the comment.", ""]
    for folder in dict.fromkeys(m.folder for m in MEASURES):
        lines.append(f"// ===== {folder} =====")
        for m in (m for m in MEASURES if m.folder == folder):
            lines.append(f"// {m.description}  |  home table: {m.table}  |  format: {m.fmt.replace(chr(92), '')}")
            body = m.expression.replace("\n", "\n    ")
            lines.append(f"{m.name} =\n    {body}\n")
    (PBI / "measures.dax").write_text("\n".join(lines))
    print(f"  measures.dax  {len(MEASURES)} measures")


def write_measures_query() -> None:
    """DAX query for Power BI Desktop's DAX query view: defines every measure, then returns the tie-out
    KPIs. Run it, compare with BUILD_GUIDE.md, then click "Update model with changes" to add the measures."""
    kpis = ('"Revenue", [Revenue], "Orders", [Orders], "Customers", [Customers], '
            '"Avg Order Value", [Avg Order Value], "Cancellation Rate", [Cancellation Rate]')
    lines = ["// Power BI Desktop: DAX query view > paste all of this > Run.",
             "// Check the result against the tie-out table in BUILD_GUIDE.md, then click",
             "// \"Update model with changes\" to add all measures to the model in one go.",
             "DEFINE"]
    for m in MEASURES:
        body = m.expression.replace("\n", "\n        ")
        lines.append(f"    MEASURE {m.table}[{m.name}] =\n        {body}")
    lines += [
        "",
        "EVALUATE",
        "UNION(",
        f'    ROW("Slice", "All data", {kpis}),',
        f'    CALCULATETABLE(ROW("Slice", "Calendar 2011", {kpis}), dim_date[year] = 2011),',
        f'    CALCULATETABLE(ROW("Slice", "France, all dates", {kpis}), dim_country[country] = "France"),',
        f'    CALCULATETABLE(ROW("Slice", "Germany, Q1 2011", {kpis}), dim_country[country] = "Germany",',
        "        DATESBETWEEN(dim_date[date], DATE(2011, 1, 1), DATE(2011, 3, 31)))",
        ")",
        "",
    ]
    (PBI / "measures_query.dax").write_text("\n".join(lines))
    print(f"  measures_query.dax  {len(MEASURES)} measures + tie-out query")


def tieout(df: pd.DataFrame) -> str:
    """Reference numbers from retail/metrics.py for checking the Power BI build."""
    lo, hi = df["invoice_date"].min().date(), df["invoice_date"].max().date()
    cases = [
        ("All data", lo, hi, ()),
        ("Calendar 2011", date(2011, 1, 1), hi, ()),
        ("France, all dates", lo, hi, ("France",)),
        ("Germany, Q1 2011", date(2011, 1, 1), date(2011, 3, 31), ("Germany",)),
    ]
    rows = ["| Slice | Revenue | Orders | Customers | Avg Order Value | Cancellation Rate |",
            "|---|---:|---:|---:|---:|---:|"]
    for label, start, end, countries in cases:
        k = metrics.kpis(metrics.filter_frame(df, start, end, countries))
        rows.append(f"| {label} | £{k['revenue']:,.2f} | {k['orders']:,} | {k['customers']:,} | "
                    f"£{k['aov']:,.2f} | {k['cancel_rate']:.2%} |")
    seg = metrics.segment_summary(metrics.rfm(df, hi + timedelta(days=1))).set_index("segment")
    cohorts = metrics.cohort_table(df, metrics.first_purchase_month(df)).set_index(["cohort", "period"])
    rules = metrics.basket_rules(df, min_support=BASKET_MIN_SUPPORT)
    top = metrics.top_products(df, "revenue", 1).iloc[0]
    extra = [
        "",
        f"- **Top product by revenue:** {top['description']} ({top['stock_code']}), £{top['revenue']:,.2f}",
        f"- **Champions:** {int(seg.loc['Champions', 'customers']):,} customers, "
        f"{seg.loc['Champions', 'revenue_share']:.2%} of segmented revenue",
        f"- **Jan 2010 cohort retention, month 1:** "
        f"{cohorts.loc[(pd.Timestamp('2010-01-01'), 1), 'retention']:.2%}",
        f"- **Basket rules at {BASKET_MIN_SUPPORT:.0%} support:** {len(rules):,} rules; strongest lift "
        f"{rules['lift'].iloc[0]:.2f} ({rules['antecedent_desc'].iloc[0]} → {rules['consequent_desc'].iloc[0]})",
    ]
    return "\n".join(rows + extra)


def refresh_guide(block: str) -> None:
    for guide in (PBI / "BUILD_GUIDE.md", PBI / "DESKTOP_GUIDE.md"):
        text = guide.read_text()
        new = re.sub(r"(<!-- TIEOUT:START -->\n).*?(\n<!-- TIEOUT:END -->)",
                     lambda m: m.group(1) + block + m.group(2), text, flags=re.S)
        guide.write_text(new)
        print(f"  {guide.name} tie-out {'refreshed' if new != text else 'unchanged'}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--no-xlsx", action="store_true", help="skip the Excel workbook")
    args = parser.parse_args()

    df = pd.read_parquet(ROOT / "data" / "processed" / "transactions.parquet")
    frames = build_frames(df)
    print("CSV files")
    write_csvs(frames)
    if not args.no_xlsx:
        print("Excel workbook (slow)")
        write_excel(frames)
    write_measures_dax()
    write_measures_query()

    import pbip  # scripts/pbip.py
    tables = model_tables()
    for t in tables:
        t.frame = frames[t.name]
    pbip.write_project(PBI, tables, RELATIONSHIPS, MEASURES)
    refresh_guide(tieout(df))


if __name__ == "__main__":
    main()
