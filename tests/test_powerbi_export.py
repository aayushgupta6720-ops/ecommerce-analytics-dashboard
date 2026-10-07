"""The Power BI export must be internally consistent and agree with the app's metrics."""

import sys
from pathlib import Path

import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

import export_powerbi as ex  # noqa: E402
from retail import metrics  # noqa: E402


@pytest.fixture(scope="module")
def source() -> pd.DataFrame:
    return pd.read_parquet(ROOT / "data" / "processed" / "transactions.parquet")


@pytest.fixture(scope="module")
def frames(source) -> dict[str, pd.DataFrame]:
    return ex.build_frames(source)


def test_frames_match_the_model_definition(frames):
    for table in ex.model_tables():
        assert list(frames[table.name].columns) == [c.name for c in table.columns], table.name


@pytest.mark.parametrize("fact_table, fact_col, dim, key", ex.RELATIONSHIPS)
def test_every_foreign_key_exists_in_its_dimension(frames, fact_table, fact_col, dim, key):
    keys = frames[dim][key]
    assert keys.is_unique, f"{dim}[{key}] must be unique"
    values = frames[fact_table][fact_col].dropna()
    missing = set(values.unique()) - set(keys)
    assert not missing, f"{len(missing)} {fact_table}[{fact_col}] values missing from {dim}, e.g. {list(missing)[:3]}"


def test_fact_fits_one_excel_sheet(frames):
    assert len(frames["fact_sales"]) <= ex.EXCEL_MAX_ROWS


def test_fact_totals_equal_app_kpis(frames, source):
    k = metrics.kpis(source)
    fact = frames["fact_sales"]
    sales, cancels = fact[~fact["is_cancellation"]], fact[fact["is_cancellation"]]
    # DAX [Revenue] = SUM(revenue) over every line (net); [Gross Sales] = sales lines only.
    assert fact["revenue"].sum() == pytest.approx(k["revenue"], abs=0.01)
    assert sales["revenue"].sum() == pytest.approx(k["gross_sales"], abs=0.01)
    assert sales["invoice"].nunique() == k["orders"]
    assert sales["customer_id"].nunique() == k["customers"]
    assert fact["quantity"].sum() == k["units"]
    assert -cancels["revenue"].sum() == pytest.approx(k["cancelled_value"], abs=0.01)


def test_date_dimension_is_continuous(frames):
    dates = frames["dim_date"]["date"]
    assert (dates.diff().dropna() == pd.Timedelta(days=1)).all()
    assert dates.min() <= frames["fact_sales"]["invoice_date"].min()
    assert dates.max() >= frames["fact_sales"]["invoice_date"].max()


def test_every_measure_has_a_home_table():
    names = {t.name for t in ex.model_tables()}
    assert all(m.table in names for m in ex.MEASURES)
    assert len({m.name for m in ex.MEASURES}) == len(ex.MEASURES)


@pytest.mark.skipif(not (ROOT / "powerbi" / "RetailAnalytics.pbip").exists(), reason="run export_powerbi.py first")
@pytest.mark.skipif(not (ROOT / ".cache" / "json-schemas" / "fabric").exists(), reason="schemas not fetched")
def test_generated_project_validates():
    import validate_pbip
    assert validate_pbip.main() == 0


@pytest.mark.skipif(not (ROOT / "powerbi" / "data" / "RetailAnalytics.xlsx").exists(), reason="workbook not exported")
def test_workbook_round_trips(frames):
    """Every sheet has all its columns filled (guards against the column-order write bug)."""
    from openpyxl import load_workbook
    wb = load_workbook(ROOT / "powerbi" / "data" / "RetailAnalytics.xlsx", read_only=True)
    for name, frame in frames.items():
        ws = wb[name]
        header, first = [row for row in ws.iter_rows(min_row=1, max_row=2, values_only=True)]
        assert list(header) == list(frame.columns), name
        expected = frame.iloc[0]
        for col, got in zip(frame.columns, first):
            want = expected[col]
            if pd.isna(want):
                assert got is None, (name, col)
            elif isinstance(want, pd.Timestamp):
                assert pd.Timestamp(got) == want, (name, col)
            elif isinstance(want, (float, int)) and not isinstance(want, bool):
                assert got == pytest.approx(want), (name, col)
            else:
                assert got == want, (name, col, got, want)


@pytest.mark.skipif(not (ROOT / "powerbi" / "data" / "RetailAnalytics.xlsx").exists(), reason="workbook not exported")
def test_workbook_sheets_are_excel_tables(frames):
    """The Power BI service only imports data formatted as Excel tables; names must match the model."""
    import re
    import zipfile
    from openpyxl.utils import get_column_letter
    # Read the table definitions straight from the xlsx; loading 1M rows through openpyxl is slow.
    with zipfile.ZipFile(ROOT / "powerbi" / "data" / "RetailAnalytics.xlsx") as z:
        tables = {}
        for n in z.namelist():
            if n.startswith("xl/tables/"):
                xml = z.read(n).decode()
                tables[re.search(r'displayName="([^"]+)"', xml).group(1)] = re.search(r' ref="([^"]+)"', xml).group(1)
    for name, frame in frames.items():
        assert tables.get(name) == f"A1:{get_column_letter(frame.shape[1])}{len(frame) + 1}", name


@pytest.mark.skipif(not (ROOT / "powerbi" / "measures_query.dax").exists(), reason="run export_powerbi.py first")
def test_measures_query_defines_every_measure():
    q = (ROOT / "powerbi" / "measures_query.dax").read_text()
    for m in ex.MEASURES:
        assert f"MEASURE {m.table}[{m.name}] =" in q, m.name
    assert q.count("EVALUATE") == 1 and q.count("ROW(") == 4
